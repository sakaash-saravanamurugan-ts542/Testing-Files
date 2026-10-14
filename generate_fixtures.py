"""Generates the fixture files used by docs/xls-read-support-testcases.md.

Usage: python3 generate_fixtures.py <output-dir>

Requires: openpyxl, pillow, xlwt, msoffcrypto-tool.

Note: the "Place in Cell" picture (Pictures!F6) is built by hand-writing the richData parts that
SheetUtil.extractInCellImages() reads. It matches the structure Excel writes, but re-save the file
in Excel if you need a byte-for-byte genuine one.
"""
import datetime
import io
import os
import random
import re
import shutil
import sys
import zipfile

import msoffcrypto
import xlwt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image

MAX_5MB = 5 * 1024 * 1024


def png_bytes(color, size=(40, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def jpeg_bytes(color, size=(40, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "JPEG")
    return buf.getvalue()


def add_image(ws, data, anchor):
    ws.add_image(XlImage(io.BytesIO(data)), anchor)


def fill_features(wb):
    ws = wb.active
    ws.title = "Sales"
    for row in [["Region", "Product", "Amount"], ["North", "Pen", 10], ["South", "Book", 25], ["East", "Bag", 7.5]]:
        ws.append(row)

    ws = wb.create_sheet("Types")
    ws.append(["Name", "Date", "Time", "Stamp", "Pct", "Flag", "Int", "Dec", "Err"])
    ws.append(["a", datetime.date(2024, 3, 15), datetime.time(14, 30, 0),
               datetime.datetime(2024, 3, 15, 14, 30, 0), 0.25, True, 4, 3.14, "#DIV/0!"])
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["C2"].number_format = "hh:mm:ss"
    ws["D2"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["E2"].number_format = "0%"
    # Row 3: extreme numbers (testcase 44)
    ws.append([None, 9007199254740993, -12.5, 1e300, 0.30000000000000004, -0.0, 0])
    # Row 4: format edge cases (testcase 46) and percent variants (testcase 47)
    ws.append([None, 0.6041666667, 1.5, datetime.date(2024, 3, 15), 4, datetime.date(2024, 3, 15), 0.25, 0.25])
    ws["B4"].number_format = "h:mm AM/PM"
    ws["C4"].number_format = "[h]:mm:ss"
    ws["D4"].number_format = 'yyyy"y"'
    ws["E4"].number_format = '0"d"'
    ws["F4"].number_format = "mmm-yy"
    ws["G4"].number_format = "0.00%"
    ws["H4"].number_format = '0"%"'

    ws = wb.create_sheet("Links")
    ws["A1"] = "Site"
    ws["A2"] = "Zoho"
    ws["A2"].hyperlink = "https://www.zoho.com"
    ws["A3"] = "Mail"
    ws["A3"].hyperlink = "mailto:test@example.com"
    ws["A4"] = "Plain"

    # Rows 1-6 each get a label in column A so they are not skipped as blank and
    # sheet row N is list index N-1 (testcases 56-61).
    ws = wb.create_sheet("Pictures")
    for r in range(1, 7):
        ws.cell(row=r, column=1, value=f"r{r}")
    add_image(ws, png_bytes("red"), "B2")      # one floating picture
    add_image(ws, png_bytes("green"), "D4")    # two floating pictures: PNG then JPEG
    add_image(ws, jpeg_bytes("blue"), "D4")
    add_image(ws, png_bytes("orange"), "A20")  # picture-only row
    ws["F6"] = "#VALUE!"                       # placeholder cell, turned into a place-in-cell picture later

    ws = wb.create_sheet("Gaps")
    ws.append(["ID", "Name"])
    ws.append([1, "one"])
    ws.append([2, "two"])
    ws.append([])
    ws.append([3, "three"])
    ws.append([])
    ws.append([4, "four"])

    ws = wb.create_sheet("Dupes")
    ws.append(["Name", "Name", None, "Name", None])
    ws.append([1, 2, 3, 4, 5])

    wb.create_sheet("Empty")

    ws = wb.create_sheet("Formulas")
    ws["A1"] = "=1+2"
    ws["B2"] = 5
    ws["A2"] = "=B2*2"
    ws["A3"] = "=NOSUCHFUNC(1)"
    ws["A4"] = "=SUM(Sales!C2:C4)"

    ws = wb.create_sheet("Wide")
    for r in range(1, 61):
        ws.append([f"R{r}C{c}" if r == 1 else r * 100 + c for c in range(1, 27)])


RICH_DATA_PARTS = {
    "xl/metadata.xml": (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
        '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000"/></metadataTypes>'
        '<futureMetadata name="XLRICHVALUE" count="1"><bk><extLst>'
        '<ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}"><xlrd:rvb i="0"/></ext>'
        '</extLst></bk></futureMetadata>'
        '<valueMetadata count="1"><bk><rc t="1" v="0"/></bk></valueMetadata>'
        '</metadata>'),
    "xl/richData/rdrichvaluestructure.xml": (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1">'
        '<s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s>'
        '</rvStructures>'),
    "xl/richData/rdrichvalue.xml": (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1">'
        '<rv s="0"><v>0</v><v>5</v></rv></rvData>'),
    "xl/richData/richValueRel.xml": (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<rel r:id="rId1"/></richValueRels>'),
    "xl/richData/_rels/richValueRel.xml.rels": (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        'Target="../media/incell1.png"/></Relationships>'),
}

CONTENT_TYPE_OVERRIDES = (
    '<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>'
    '<Override PartName="/xl/richData/rdrichvalue.xml" ContentType="application/vnd.ms-excel.rdrichvalue+xml"/>'
    '<Override PartName="/xl/richData/rdrichvaluestructure.xml" ContentType="application/vnd.ms-excel.rdrichvaluestructure+xml"/>'
    '<Override PartName="/xl/richData/richValueRel.xml" ContentType="application/vnd.ms-excel.richvaluerel+xml"/>')

WORKBOOK_RELS = (
    '<Relationship Id="rIdMeta" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" Target="metadata.xml"/>'
    '<Relationship Id="rIdRv" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValue" Target="richData/rdrichvalue.xml"/>'
    '<Relationship Id="rIdRvs" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueStructure" Target="richData/rdrichvaluestructure.xml"/>'
    '<Relationship Id="rIdRvr" Type="http://schemas.microsoft.com/office/2022/10/relationships/richValueRel" Target="richData/richValueRel.xml"/>')


def add_incell_picture(src, dst, pictures_sheet_xml="xl/worksheets/sheet4.xml", overrides=None):
    """Copies src to dst, adding the richData parts and marking Pictures!F6 as an in-cell picture.

    overrides maps a part name to replacement text, used to build the broken/XXE variants.
    """
    overrides = overrides or {}
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                if 'Extension="png"' not in text:
                    text = text.replace("<Override", '<Default Extension="png" ContentType="image/png"/><Override', 1)
                data = text.replace("</Types>", CONTENT_TYPE_OVERRIDES + "</Types>").encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.decode("utf-8").replace("</Relationships>", WORKBOOK_RELS + "</Relationships>").encode("utf-8")
            elif item.filename == pictures_sheet_xml:
                text = data.decode("utf-8")
                text, n = re.subn(r'<c r="F6"', '<c r="F6" vm="1"', text)
                assert n == 1, "F6 not found in " + pictures_sheet_xml
                data = text.encode("utf-8")
            zout.writestr(item.filename, data)
        for name, text in RICH_DATA_PARTS.items():
            zout.writestr(name, overrides.get(name, text))
        zout.writestr("xl/media/incell1.png", png_bytes("purple"))


def build_features(path, workdir):
    base = os.path.join(workdir, "features_base.xlsx")
    wb = Workbook()
    fill_features(wb)
    wb.save(base)
    add_incell_picture(base, path)
    return base


def build_big(path):
    """Valid workbook larger than 5 MB (random text does not compress well)."""
    rng = random.Random(1)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sales")
    ws.append(["Region", "Product", "Amount"])
    alphabet = "abcdefghijklmnopqrstuvwxyz0123456789"
    for _ in range(250000):
        ws.append(["".join(rng.choices(alphabet, k=20)), "".join(rng.choices(alphabet, k=20)), rng.random()])
    wb.save(path)


def build_sized_pair(out, workdir):
    """Valid workbooks of exactly 5,242,880 and 5,242,881 bytes, padded with a zip comment."""
    rng = random.Random(3)
    alphabet = "abcdefghijklmnopqrstuvwxyz0123456789"
    rows = 90000
    while True:
        base = os.path.join(workdir, "sized_base.xlsx")
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Sales")
        ws.append(["Region", "Product", "Amount"])
        for _ in range(rows):
            ws.append(["".join(rng.choices(alphabet, k=20)), "".join(rng.choices(alphabet, k=20)), rng.random()])
        wb.save(base)
        size = os.path.getsize(base)
        gap = MAX_5MB - size
        if 0 < gap < 60000:
            break
        rows = int(rows * MAX_5MB / size) - 100  # land just below the target, then pad
    for name, extra in (("5MB-exact.xlsx", 0), ("5MB-plus-one.xlsx", 1)):
        dst = os.path.join(out, name)
        shutil.copyfile(base, dst)
        with zipfile.ZipFile(dst, "a") as zf:
            zf.comment = b"x" * (gap + extra)
        print(name, os.path.getsize(dst))


def build_100k(path):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["Id", "Name", "Value"])
    for i in range(1, 100001):
        ws.append([i, f"name{i % 100}", i / 7])
    wb.save(path)


def build_legacy_xls(path):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sales")
    for r, row in enumerate([["Region", "Product", "Amount"], ["North", "Pen", 10], ["South", "Book", 25], ["East", "Bag", 7.5]]):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    wb.add_sheet("Other").write(0, 0, "x")
    wb.save(path)


def build_encrypted(src, dst):
    with open(src, "rb") as f, open(dst, "wb") as out:
        msoffcrypto.OfficeFile(f).encrypt("password", out)


def write_zip(path, entries):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(name, data)


def main(out):
    os.makedirs(out, exist_ok=True)
    work = os.path.join(out, ".work")
    os.makedirs(work, exist_ok=True)
    p = lambda name: os.path.join(out, name)

    build_features(p("features.xlsx"), work)
    shutil.copyfile(p("features.xlsx"), p("a.xlsx"))
    build_encrypted(p("features.xlsx"), p("encrypted.xlsx"))
    build_legacy_xls(p("legacy.xls"))

    # xlsm is no longer supported (testcase 78): same workbook content, macro-enabled name.
    shutil.copyfile(p("features.xlsx"), p("macros.xlsm"))

    # Broken in-cell image metadata (testcase 77) and XXE attempt (testcase 76).
    base = os.path.join(work, "features_base.xlsx")
    add_incell_picture(base, p("broken_richdata.xlsx"), overrides={
        "xl/richData/rdrichvalue.xml": RICH_DATA_PARTS["xl/richData/rdrichvalue.xml"][:80]})
    add_incell_picture(base, p("xxe.xlsx"), overrides={
        "xl/metadata.xml": '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><metadata>&e;</metadata>'})

    build_big(p("big.xlsx"))
    build_sized_pair(out, work)
    build_100k(p("100k-rows.xlsx"))

    write_zip(p("bomb_entries.xlsx"), [(f"e{i}.xml", b"x") for i in range(1001)])
    write_zip(p("bomb_size.xlsx"), [("xl/big.xml", b"\0" * (60 * 1024 * 1024))])
    write_zip(p("bomb_ratio.xlsx"), [("xl/big.xml", b"\0" * (20 * 1024 * 1024))])
    shutil.copyfile(p("bomb_ratio.xlsx"), p("ok.xls"))  # renamed zip bomb (testcase 15)

    with open(p("corrupt.xlsx"), "wb") as f:
        f.write(random.Random(2).randbytes(4096))
    with open(p("notes.txt"), "w") as f:
        f.write("just a text file\n")
    with open(p("data.csv"), "w") as f:
        f.write("Region,Product,Amount\nNorth,Pen,10\n")
    with open(p("report.pdf"), "wb") as f:
        f.write(b"%PDF-1.1\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
                b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n")

    shutil.rmtree(work)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "fixtures")
